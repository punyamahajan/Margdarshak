import uuid
import csv
from io import BytesIO, StringIO
from collections import Counter
from datetime import datetime, timezone
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.session import get_db_session
from app.core.security import require_coordinator
from app.models.admin import KnowledgeDocument, ShortlistRecord, StudentNotification, TicketCluster, TicketWorkflow
from app.models.placement_drive import PlacementDrive
from app.models.student import Student
from app.models.ticket import Ticket
from app.models.case_card import CaseCard
from app.services.admin_workflow_service import calculate_priority
from app.services.agora_service import AgoraServiceError, generate_rtc_token
from app.core.config import get_settings

router = APIRouter(prefix="/admin", tags=["admin"], dependencies=[Depends(require_coordinator)])


class WorkflowUpdate(BaseModel):
    status: str | None = Field(default=None, pattern="^(open|claimed|waiting|escalated|resolved)$")
    assigned_coordinator: str | None = None
    urgency: str | None = Field(default=None, pattern="^(low|medium|high|critical)$")
    language: str | None = Field(default=None, max_length=50)
    category: str | None = Field(default=None, max_length=100)
    original_request: str | None = Field(default=None, max_length=3000)
    conversation_summary: str | None = Field(default=None, max_length=3000)
    cluster_id: uuid.UUID | None = None


class ClusterWrite(BaseModel):
    title: str = Field(min_length=2, max_length=255)
    company_name: str | None = None
    drive_id: uuid.UUID | None = None
    urgency: str = "medium"


class DraftWrite(BaseModel):
    notes: str = Field(min_length=3, max_length=3000)


class UpdateWrite(BaseModel):
    message: str = Field(min_length=3, max_length=3000)


class KnowledgeWrite(BaseModel):
    title: str = Field(min_length=2, max_length=255)
    document_type: str = "placement_policy"
    company_name: str | None = None
    drive_id: uuid.UUID | None = None
    version_label: str
    source_reference: str
    status: str = "draft"
    content: dict[str, Any] = Field(default_factory=dict)


class ShortlistImportConfirm(BaseModel):
    title: str = Field(min_length=2, max_length=255)
    source_reference: str = Field(min_length=3, max_length=2048)
    version_label: str = Field(min_length=1, max_length=100)
    rows: list[dict[str, Any]] = Field(min_length=1, max_length=5000)


REQUIRED_SHORTLIST_COLUMNS = {"name", "enrollment_number", "email", "company", "drive_id", "shortlisted", "role", "round"}


def _normalize_shortlist_row(row: dict[str, Any]) -> dict[str, Any]:
    return {str(key).strip().lower(): ("" if value is None else str(value).strip()) for key, value in row.items()}


async def _parse_shortlist(file: UploadFile) -> list[dict[str, Any]]:
    content = await file.read()
    filename = (file.filename or "").lower()
    if len(content) > 5_000_000:
        raise HTTPException(413, "file must be smaller than 5 MB")
    try:
        if filename.endswith(".csv"):
            rows = list(csv.DictReader(StringIO(content.decode("utf-8-sig"))))
        elif filename.endswith(".xlsx"):
            from openpyxl import load_workbook
            sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).active
            values = list(sheet.iter_rows(values_only=True))
            if not values: rows = []
            else: rows = [dict(zip(values[0], item)) for item in values[1:]]
        else:
            raise HTTPException(400, "upload a CSV or XLSX file")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, "we couldn't read that shortlist file") from exc
    normalized = [_normalize_shortlist_row(row) for row in rows if any(value not in (None, "") for value in row.values())]
    columns = set(normalized[0]) if normalized else set()
    missing = sorted(REQUIRED_SHORTLIST_COLUMNS - columns)
    if missing:
        raise HTTPException(422, {"message": "missing required shortlist columns", "columns": missing})
    return normalized


def _ticket_payload(ticket: Ticket, student: Student, drive: PlacementDrive | None, workflow: TicketWorkflow | None, case_card: dict[str, Any] | None = None) -> dict[str, Any]:
    context = case_card or {}
    policy = context.get("policy") if isinstance(context.get("policy"), dict) else {}
    return {
        "id": str(ticket.id), "status": workflow.status if workflow else ticket.status.value,
        "issue_summary": ticket.issue_summary or "Conversation needs review",
        "confidence_score": ticket.confidence_score, "created_at": ticket.created_at, "updated_at": workflow.updated_at if workflow else ticket.created_at,
        "student": {"name": student.name, "enrollment_number": student.roll_number, "email": student.email},
        "placement": {"company": drive.company_name if drive else policy.get("company_name"), "drive_id": str(ticket.drive_id) if ticket.drive_id else context.get("drive_id"), "role": context.get("role") or policy.get("role"), "round": context.get("round") or policy.get("round"), "deadline": context.get("deadline") or policy.get("deadline")},
        "intelligence": {"urgency": workflow.urgency if workflow else "medium", "language": workflow.language if workflow else "English", "category": workflow.category if workflow else "Other", "assigned_coordinator": workflow.assigned_coordinator if workflow else None, "cluster_id": str(workflow.cluster_id) if workflow and workflow.cluster_id else None},
        "original_request": workflow.original_request if workflow else None,
        "conversation_summary": workflow.conversation_summary if workflow else None,
        "source": context.get("source") or (drive.policy_doc_ref if drive else None),
    }


async def _workflow(db: AsyncSession, ticket_id: uuid.UUID) -> TicketWorkflow:
    workflow = await db.get(TicketWorkflow, ticket_id)
    if workflow is None:
        workflow = TicketWorkflow(ticket_id=ticket_id)
        db.add(workflow)
        await db.flush()
    return workflow


@router.get("/tickets")
async def list_admin_tickets(
    company: str | None = None,
    drive_id: uuid.UUID | None = None,
    coordinator: str | None = None,
    language: str | None = None,
    urgency: str | None = None,
    round_name: str | None = None,
    status: str | None = Query(default=None, pattern="^(open|claimed|waiting|escalated|resolved)$"),
    db: AsyncSession = Depends(get_db_session),
) -> list[dict[str, Any]]:
    rows = (await db.execute(select(Ticket, Student, PlacementDrive, TicketWorkflow).join(Student, Ticket.student_id == Student.id).outerjoin(PlacementDrive, Ticket.drive_id == PlacementDrive.id).outerjoin(TicketWorkflow, TicketWorkflow.ticket_id == Ticket.id).order_by(Ticket.created_at.desc()))).all()
    cards = (await db.scalars(select(CaseCard).order_by(CaseCard.last_updated))).all()
    card_by_ticket = {card.ticket_id: dict(card.structured_json or {}) for card in cards}
    results = [_ticket_payload(*row, card_by_ticket.get(row[0].id)) for row in rows]
    def matches(ticket: dict[str, Any]) -> bool:
        intelligence = ticket["intelligence"]
        placement = ticket["placement"]
        return (
            (company is None or (placement["company"] or "").lower() == company.lower())
            and (drive_id is None or placement["drive_id"] == str(drive_id))
            and (coordinator is None or (intelligence["assigned_coordinator"] or "").lower() == coordinator.lower())
            and (language is None or intelligence["language"].lower() == language.lower())
            and (urgency is None or intelligence["urgency"] == urgency.lower())
            and (round_name is None or (placement["round"] or "").lower() == round_name.lower())
            and (status is None or ticket["status"] == status)
        )
    return [ticket for ticket in results if matches(ticket)]


@router.patch("/tickets/{ticket_id}")
async def update_ticket_workflow(ticket_id: uuid.UUID, payload: WorkflowUpdate, db: AsyncSession = Depends(get_db_session)) -> dict[str, Any]:
    if await db.get(Ticket, ticket_id) is None:
        raise HTTPException(404, "ticket not found")
    if payload.cluster_id is not None and await db.get(TicketCluster, payload.cluster_id) is None:
        raise HTTPException(404, "cluster not found")
    workflow = await _workflow(db, ticket_id)
    for key, value in payload.model_dump(exclude_none=True).items(): setattr(workflow, key, value)
    await db.commit()
    return {"ticket_id": str(ticket_id), "status": workflow.status, "updated_at": workflow.updated_at}


@router.get("/overview")
async def overview(db: AsyncSession = Depends(get_db_session)) -> dict[str, Any]:
    tickets = await list_admin_tickets(db=db)
    counts = Counter(ticket["status"] for ticket in tickets)
    drives = (await db.scalars(select(PlacementDrive).where(PlacementDrive.status == "active"))).all()
    priority = sorted((ticket for ticket in tickets if ticket["status"] != "resolved"), key=lambda item: ({"critical": 4, "high": 3, "medium": 2, "low": 1}.get(item["intelligence"]["urgency"], 0), item["created_at"]), reverse=True)[:3]
    knowledge = (await db.scalars(select(KnowledgeDocument).order_by(KnowledgeDocument.created_at.desc()).limit(5))).all()
    recent_updates = ([{"kind": "ticket", "text": f"{item['student']['enrollment_number']}: {item['issue_summary']}", "at": item["updated_at"]} for item in tickets[:5]] + [{"kind": "knowledge", "text": f"{doc.title} is {doc.status}", "at": doc.created_at} for doc in knowledge])
    recent_updates.sort(key=lambda item: item["at"], reverse=True)
    return {"metrics": {key: counts.get(key, 0) for key in ("open", "claimed", "waiting", "escalated", "resolved")}, "priority_tickets": priority, "active_drives": [{"id": str(drive.id), "company": drive.company_name, "status": drive.status, "policy": drive.policy_doc_ref} for drive in drives], "recent_updates": recent_updates[:6]}


@router.get("/stats")
async def admin_stats(db: AsyncSession = Depends(get_db_session)) -> dict[str, Any]:
    tickets = await list_admin_tickets(db=db)
    category_counts = Counter(item["intelligence"]["category"] for item in tickets)
    language_counts = Counter(item["intelligence"]["language"] for item in tickets)
    company_counts = Counter(item["placement"]["company"] or "Unassigned" for item in tickets)
    clusters = await list_clusters(db)
    resolved_hours = [max(0.0, (item["updated_at"] - item["created_at"]).total_seconds() / 3600) for item in tickets if item["status"] == "resolved"]
    return {
        "total_conversations": len(tickets),
        "tickets_created": len(tickets),
        "tickets_resolved": sum(item["status"] == "resolved" for item in tickets),
        "students_assisted": len({item["student"]["enrollment_number"] for item in tickets}),
        "active_clusters": sum(item["status"] != "resolved" for item in clusters),
        "average_resolution_time_hours": round(sum(resolved_hours) / len(resolved_hours), 1) if resolved_hours else 0.0,
        "most_common_issue_categories": dict(category_counts.most_common(10)),
        "most_requested_companies": dict(company_counts.most_common(10)),
        "most_common_languages": dict(language_counts.most_common(10)),
    }


@router.get("/clusters")
async def list_clusters(db: AsyncSession = Depends(get_db_session)) -> list[dict[str, Any]]:
    clusters = (await db.scalars(select(TicketCluster).order_by(TicketCluster.created_at.desc()))).all()
    results = []
    for cluster in clusters:
        workflows = (await db.scalars(select(TicketWorkflow).where(TicketWorkflow.cluster_id == cluster.id))).all()
        documents = (await db.scalars(select(KnowledgeDocument).where(KnowledgeDocument.company_name == cluster.company_name, KnowledgeDocument.status == "published"))).all()
        deadline_days: int | None = None
        for document in documents:
            raw_deadline = (document.content or {}).get("deadline")
            if raw_deadline:
                try:
                    days = (datetime.fromisoformat(str(raw_deadline)).date() - datetime.now(timezone.utc).date()).days
                    deadline_days = days if deadline_days is None else min(deadline_days, days)
                except ValueError:
                    continue
        score, priority = calculate_priority(len(workflows), [item.urgency for item in workflows], blocked="link" in cluster.title.lower(), deadline_days=deadline_days)
        results.append({"id": str(cluster.id), "title": cluster.title, "company_name": cluster.company_name, "urgency": priority, "priority_score": score, "status": cluster.status, "affected_students": len(workflows), "incident_update": cluster.incident_update, "response_draft": cluster.response_draft, "created_at": cluster.created_at})
    return results


@router.post("/clusters")
async def create_cluster(payload: ClusterWrite, db: AsyncSession = Depends(get_db_session)) -> dict[str, Any]:
    cluster = TicketCluster(**payload.model_dump())
    db.add(cluster); await db.commit(); await db.refresh(cluster)
    return {"id": str(cluster.id)}


@router.post("/agora/session")
async def start_admin_agora_session() -> dict[str, Any]:
    """Issue a short-lived RTC credential for coordinator voice drafting."""
    channel = f"admin-draft-{uuid.uuid4().hex}"
    uid = 2
    try:
        token = generate_rtc_token(channel, uid)
    except AgoraServiceError as exc:
        raise HTTPException(503, "Agora voice drafting is not configured") from exc
    return {"app_id": get_settings().agora_app_id, "channel_name": channel, "uid": uid, "rtc_token": token}


@router.post("/clusters/{cluster_id}/draft-response")
async def draft_cluster_response(cluster_id: uuid.UUID, payload: DraftWrite, db: AsyncSession = Depends(get_db_session)) -> dict[str, str]:
    cluster = await db.get(TicketCluster, cluster_id)
    if cluster is None: raise HTTPException(404, "cluster not found")
    cluster.response_draft = f"PLACEMENT UPDATE\n\n{payload.notes.strip()}\n\nIf the issue continues, please raise a new query through Margdarshak."
    await db.commit(); return {"draft": cluster.response_draft}


@router.post("/clusters/{cluster_id}/publish-update")
async def publish_cluster_update(cluster_id: uuid.UUID, payload: UpdateWrite, db: AsyncSession = Depends(get_db_session)) -> dict[str, Any]:
    cluster = await db.get(TicketCluster, cluster_id)
    if cluster is None: raise HTTPException(404, "cluster not found")
    cluster.incident_update = payload.message
    tickets = (await db.scalars(select(Ticket).join(TicketWorkflow, TicketWorkflow.ticket_id == Ticket.id).where(TicketWorkflow.cluster_id == cluster_id))).all()
    db.add_all([StudentNotification(student_id=ticket.student_id, ticket_id=ticket.id, title="Placement update", body=payload.message) for ticket in tickets])
    await db.commit(); return {"published_at": datetime.now(timezone.utc), "cluster_id": str(cluster.id), "notified_students": len(tickets)}


@router.post("/clusters/{cluster_id}/resolve")
async def resolve_cluster(cluster_id: uuid.UUID, db: AsyncSession = Depends(get_db_session)) -> dict[str, Any]:
    cluster = await db.get(TicketCluster, cluster_id)
    if cluster is None: raise HTTPException(404, "cluster not found")
    cluster.status = "resolved"
    workflows = (await db.scalars(select(TicketWorkflow).where(TicketWorkflow.cluster_id == cluster_id))).all()
    for workflow in workflows: workflow.status = "resolved"
    db.add(KnowledgeDocument(title=f"Resolved issue: {cluster.title}", document_type="resolved_issue", company_name=cluster.company_name, version_label=datetime.now().strftime("%Y.%m"), status="published", source_reference="coordinator://cluster-resolution", content={"resolution": cluster.incident_update or cluster.response_draft or "Resolved by coordinator."}))
    await db.commit(); return {"resolved": len(workflows), "cluster_id": str(cluster.id)}


@router.get("/knowledge")
async def list_knowledge(db: AsyncSession = Depends(get_db_session)) -> list[dict[str, Any]]:
    documents = (await db.scalars(select(KnowledgeDocument).order_by(KnowledgeDocument.created_at.desc()))).all()
    return [{"id": str(doc.id), "title": doc.title, "type": doc.document_type, "company": doc.company_name, "version": doc.version_label, "status": doc.status, "source": doc.source_reference, "content": doc.content, "created_at": doc.created_at} for doc in documents]


class KnowledgeStatusUpdate(BaseModel):
    status: str | None = Field(default=None, pattern="^(draft|published|expired)$")


@router.post("/knowledge")
@router.post("/knowledge/policy")
async def create_knowledge(payload: KnowledgeWrite, db: AsyncSession = Depends(get_db_session)) -> dict[str, str]:
    document = KnowledgeDocument(**payload.model_dump())
    db.add(document)
    await db.commit()
    return {"id": str(document.id), "status": document.status}


@router.post("/knowledge/shortlist-preview")
@router.post("/knowledge/preview")
async def shortlist_preview(file: UploadFile = File(...)) -> dict[str, Any]:
    rows = await _parse_shortlist(file)
    return {"valid": True, "record_count": len(rows), "preview": rows[:20], "rows": rows}


@router.post("/knowledge/shortlist-import")
@router.post("/knowledge/import")
async def shortlist_import(payload: ShortlistImportConfirm, db: AsyncSession = Depends(get_db_session)) -> dict[str, Any]:
    document = KnowledgeDocument(title=payload.title, document_type="shortlist", version_label=payload.version_label, status="published", source_reference=payload.source_reference, content={"record_count": len(payload.rows)})
    db.add(document); await db.flush()
    records: list[ShortlistRecord] = []
    for raw in payload.rows:
        row = _normalize_shortlist_row(raw)
        if not REQUIRED_SHORTLIST_COLUMNS.issubset(row):
            raise HTTPException(422, "shortlist rows do not match the validated template")
        records.append(ShortlistRecord(knowledge_document_id=document.id, student_name=row["name"], enrollment_number=row["enrollment_number"], email=row["email"], company_name=row["company"], drive_id=row["drive_id"], shortlisted=row["shortlisted"].lower() in {"true", "yes", "1", "shortlisted"}, role=row["role"] or None, round=row["round"] or None))
    db.add_all(records); await db.commit()
    return {"document_id": str(document.id), "imported": len(records), "status": "published"}


@router.patch("/knowledge/{document_id}")
async def update_knowledge(
    document_id: uuid.UUID,
    status: str | None = Query(default=None),
    payload: KnowledgeStatusUpdate | None = None,
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, str]:
    raw_status = status if isinstance(status, str) else None
    body_status = payload.status if payload else None
    effective_status = raw_status or body_status
    if not effective_status or effective_status not in {"draft", "published", "expired"}:
        raise HTTPException(400, "invalid knowledge status")
    document = await db.get(KnowledgeDocument, document_id)
    if document is None:
        raise HTTPException(404, "knowledge document not found")
    document.status = effective_status
    await db.commit()
    return {"id": str(document.id), "status": document.status}


knowledge_router = APIRouter(prefix="/knowledge", tags=["knowledge"], dependencies=[Depends(require_coordinator)])


@knowledge_router.get("")
@knowledge_router.get("/")
async def list_knowledge_spec(db: AsyncSession = Depends(get_db_session)) -> list[dict[str, Any]]:
    return await list_knowledge(db=db)


@knowledge_router.post("")
@knowledge_router.post("/")
@knowledge_router.post("/policy")
async def create_knowledge_spec(payload: KnowledgeWrite, db: AsyncSession = Depends(get_db_session)) -> dict[str, str]:
    return await create_knowledge(payload=payload, db=db)


@knowledge_router.patch("/{document_id}")
async def update_knowledge_spec(
    document_id: uuid.UUID,
    status: str | None = Query(default=None),
    payload: KnowledgeStatusUpdate | None = None,
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, str]:
    return await update_knowledge(document_id=document_id, status=status, payload=payload, db=db)


@knowledge_router.post("/shortlist-preview")
@knowledge_router.post("/preview")
async def shortlist_preview_spec(file: UploadFile = File(...)) -> dict[str, Any]:
    return await shortlist_preview(file=file)


@knowledge_router.post("/shortlist-import")
@knowledge_router.post("/import")
async def shortlist_import_spec(payload: ShortlistImportConfirm, db: AsyncSession = Depends(get_db_session)) -> dict[str, Any]:
    return await shortlist_import(payload=payload, db=db)


stats_router = APIRouter(prefix="/stats", tags=["stats"], dependencies=[Depends(require_coordinator)])


@stats_router.get("")
@stats_router.get("/")
async def admin_stats_spec(db: AsyncSession = Depends(get_db_session)) -> dict[str, Any]:
    return await admin_stats(db=db)


